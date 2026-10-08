"""
生成《评论有用性分析系统 — 部署说明》Word 文档，保存到桌面
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DESKTOP = r"C:\Users\86183\Desktop"
OUT = os.path.join(DESKTOP, "评论有用性分析系统_部署说明.docx")

doc = Document()

# ---- 设置中文默认字体 ----
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_title(text, level=0):
    if level == 0:
        h = doc.add_heading('', level=0)
        run = h.add_run(text)
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.size = Pt(22)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        h = doc.add_heading('', level=level)
        run = h.add_run(text)
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        size_map = {1: 16, 2: 14, 3: 12}
        run.font.size = Pt(size_map.get(level, 12))
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x68)

def add_para(text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(size)
    run.bold = bold
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    return p

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(11)
    return p

def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    run.font.size = Pt(10)
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.line_spacing = 1.2
    # 灰色底
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F2F2F2')
    p._p.get_or_add_pPr().append(shd)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.size = Pt(11)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx+1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ''
            run = cells[c_idx].paragraphs[0].add_run(str(val))
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(10.5)
    return table


# ====================== 正文开始 ======================
add_title("评论有用性分析系统")
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("— 系统部署与运行说明 —")
r.font.name = '黑体'
r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
ri = info.add_run("版本 v3 · 基于 Python + Flask + 机器学习融合模型")
ri.font.size = Pt(11)
ri.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
doc.add_paragraph()

# ---------------- 1. 系统概述 ----------------
add_title("一、系统概述", 1)
add_para(
    "本系统为基于机器学习的电商评论有用性分析平台，使用随机森林、逻辑回归、"
    "LightGBM 三种基模型加权融合，对评论文本及其元数据进行有用性二分类预测；"
    "并通过 Flask + ECharts 提供可视化大屏，支持模型对比、混淆矩阵、ROC/PR 曲线、"
    "概率分布、实时单条预测、批量 CSV 上传/导出等功能。"
)
add_title("1.1 系统组成", 2)
add_bullet("数据预处理模块：data_preprocessor.py（清洗、特征工程、TF-IDF、标准化）")
add_bullet("模型训练模块：train.py、model_trainer.py、model_rf.py、model_lr.py、model_lgb.py")
add_bullet("主启动入口：main.py（自动检测模型，缺失则训练）")
add_bullet("Web 可视化大屏：web/app.py（Flask 后端） + web/templates/index.html + web/static/")
add_bullet("一键启动脚本：启动.bat（训练）、启动可视化大屏.bat（启动 Web）")
add_bullet("已训练模型：models/ 目录下 *.pkl 文件")

# ---------------- 2. 运行环境 ----------------
add_title("二、运行环境要求", 1)
add_title("2.1 硬件要求", 2)
add_table(
    ["配置项", "最低要求", "推荐配置"],
    [
        ["CPU", "Intel i5 / 同等性能双核", "Intel i7 / Ryzen 7 及以上"],
        ["内存", "8 GB", "16 GB 及以上"],
        ["硬盘", "可用空间 2 GB", "SSD 5 GB 以上"],
        ["显示器", "1366×768", "1920×1080 或更高（大屏推荐 2K）"],
        ["网络", "本地运行可离线", "首次安装依赖需联网"],
    ]
)
add_title("2.2 软件环境", 2)
add_table(
    ["软件", "版本要求", "说明"],
    [
        ["操作系统", "Windows 10 / 11（64 位）", "亦兼容 macOS、Linux"],
        ["Python", "3.8 — 3.11", "推荐 3.10，需加入系统 PATH"],
        ["pip", "≥ 21.0", "随 Python 安装"],
        ["浏览器", "Chrome 90+ / Edge 90+", "用于访问可视化大屏"],
    ]
)
add_title("2.3 Python 依赖库", 2)
add_para("依赖已统一记录在 requirements.txt 中，主要包含：")
deps = [
    ["pandas",       "2.0.3",  "数据读取与处理"],
    ["numpy",        "1.24.3", "数值计算基础库"],
    ["scikit-learn", "1.3.0",  "随机森林、逻辑回归、评估指标"],
    ["xgboost",      "2.0.3",  "梯度提升模型（备用）"],
    ["lightgbm",     "4.1.0",  "LightGBM 模型，融合主力"],
    ["jieba",        "0.42.1", "中文分词，用于 TF-IDF 特征"],
    ["openpyxl",     "3.1.2",  "读写 Excel 文件"],
    ["joblib",       "1.3.2",  "模型序列化与加载"],
    ["flask",        "3.0.0",  "Web 可视化大屏后端"],
]
add_table(["依赖包", "版本", "用途"], deps)
add_para("可选依赖：imbalanced-learn（用于 SMOTE 过采样，若缺失则自动跳过）。")

# ---------------- 3. 部署步骤 ----------------
add_title("三、部署步骤", 1)
add_title("3.1 获取源码", 2)
add_para("将完整项目文件夹（含 config.py、main.py、train.py、web/、models/ 等）"
         "拷贝至目标路径，例如 D:\\评论有用性分析毕设代码。路径建议不含空格。")

add_title("3.2 安装 Python", 2)
add_bullet("访问 https://www.python.org/downloads/ 下载 Python 3.8—3.11 安装包；")
add_bullet("安装时务必勾选 “Add Python to PATH”；")
add_bullet("安装完成后在命令行执行下列命令验证：")
add_code("python --version\npip --version")

add_title("3.3 安装依赖库", 2)
add_para("在项目根目录打开命令行（cmd 或 PowerShell），执行：")
add_code("cd /d D:\\评论有用性分析毕设代码\npip install -r requirements.txt")
add_para("如国内网络较慢，可指定清华镜像：")
add_code("pip install -r requirements.txt -i https://pypi.tuna.tsinghua.edu.cn/simple")
add_para("如需启用 SMOTE 类别平衡，可额外安装：")
add_code("pip install imbalanced-learn")

add_title("3.4 准备数据文件", 2)
add_bullet("默认训练数据为项目根目录下的 ecommerce_comments_relabeled.csv；")
add_bullet("路径在 config.py 中由 DATA_PATH 变量控制，可按需修改；")
add_bullet("如使用自定义数据，需包含字段：comment_text、useful_votes、platform、"
           "product_category、sentiment、risk_level 等（具体见 data_preprocessor.py）。")

add_title("3.5 验证模型文件", 2)
add_para("models/ 目录下应包含以下已训练文件，缺失任一文件均需重新训练：")
add_table(
    ["文件名", "用途"],
    [
        ["random_forest_model.pkl",     "随机森林基模型"],
        ["logistic_regression_model.pkl","逻辑回归基模型"],
        ["lightgbm_model.pkl",          "LightGBM 基模型"],
        ["scaler.pkl",                  "特征标准化器（StandardScaler）"],
        ["tfidf_vectorizer.pkl",        "TF-IDF 中文分词向量化器"],
        ["feature_names.pkl",           "训练特征列名"],
        ["categories.pkl",              "平台/品类/情感/风险类别字典"],
        ["best_threshold.pkl",          "融合模型最优分类阈值"],
    ]
)

# ---------------- 4. 系统启动 ----------------
add_title("四、系统启动方式", 1)

add_title("4.1 一键脚本启动（推荐）", 2)
add_para("项目根目录提供两个 .bat 一键脚本，双击即可运行：")
add_table(
    ["脚本", "功能", "适用场景"],
    [
        ["启动.bat", "检查环境 → 安装依赖 → 运行 main.py（自动训练或跳过）", "首次部署、模型重训"],
        ["启动可视化大屏.bat", "检查环境 → 启动 Flask 服务 → 自动打开浏览器", "日常使用、演示"],
    ]
)
add_para("启动可视化大屏.bat 内部执行的关键命令等价于：")
add_code("python web\\app.py")
add_para("服务启动后会监听 http://127.0.0.1:5000，浏览器自动弹出大屏页面。"
         "首次启动需加载模型并在测试集上重算指标，约 30 秒 — 1 分钟。")

add_title("4.2 命令行手动启动", 2)
add_para("（1）训练模型（如 models/ 缺失）：")
add_code("python main.py")
add_para("（2）单独运行训练脚本：")
add_code("python train.py")
add_para("（3）启动 Web 大屏：")
add_code("python web\\app.py")

add_title("4.3 验证部署成功", 2)
add_bullet("命令行无报错，输出 \"* Running on http://127.0.0.1:5000\"；")
add_bullet("浏览器访问 http://127.0.0.1:5000 能正常加载大屏页面；")
add_bullet("页面上模型对比、混淆矩阵、ROC 曲线等图表均能渲染出数据；")
add_bullet("输入一条评论文本，点击 “实时预测”，能返回有用/无用结果及概率。")

# ---------------- 5. 功能使用 ----------------
add_title("五、主要功能使用说明", 1)
add_title("5.1 模型性能监控区", 2)
add_para("展示四个模型（随机森林、逻辑回归、LightGBM、融合模型）在测试集上的"
         "准确率、精确率、召回率、F1、AUC 五项核心指标，以及混淆矩阵与 ROC、PR、"
         "F1-阈值、概率分布、校准曲线等图表，用于评估模型表现。")
add_title("5.2 实时评论预测", 2)
add_para("在输入框中粘贴或输入一条评论文本，选择平台、品类、情感、风险等元数据后，"
         "点击 “开始预测”，系统将给出四个模型各自的有用性概率与融合判定结果。")
add_title("5.3 批量数据上传与导出", 2)
add_bullet("点击 “上传 CSV”，选择符合字段要求的 CSV 文件；");
add_bullet("系统自动完成特征工程 → 四模型预测 → 结果回显到表格；")
add_bullet("如上传文件包含 useful_votes 列，将自动计算指标并刷新所有图表；")
add_bullet("可点击 “导出结果” 下载含预测概率与标签的 CSV 文件。")

# ---------------- 6. 目录结构 ----------------
add_title("六、目录结构说明", 1)
add_code(
    "评论有用性分析毕设代码/\n"
    "├── config.py                    # 全局配置（路径、模型超参、阈值、融合权重）\n"
    "├── main.py                      # 主入口：检测模型，缺失则触发训练\n"
    "├── train.py                     # 完整训练流程\n"
    "├── data_preprocessor.py         # 数据加载、特征工程、TF-IDF、标准化\n"
    "├── model_trainer.py             # 多模型训练、融合、阈值搜索、保存/加载\n"
    "├── model_rf.py / model_lr.py / model_lgb.py  # 三种基模型封装\n"
    "├── requirements.txt             # Python 依赖清单\n"
    "├── 启动.bat                     # 训练一键脚本\n"
    "├── 启动可视化大屏.bat           # Web 一键脚本\n"
    "├── ecommerce_comments_relabeled.csv  # 训练数据集\n"
    "├── batch_prediction_template.csv     # 批量预测模板\n"
    "├── models/                      # 已训练模型与预处理器\n"
    "├── outputs/                     # 训练结果报告 experiment_results.txt\n"
    "└── web/                         # 可视化大屏\n"
    "    ├── app.py                   # Flask 后端\n"
    "    ├── templates/index.html     # 大屏前端页面\n"
    "    └── static/css, static/js    # 样式与脚本"
)

# ---------------- 7. 常见问题 ----------------
add_title("七、常见问题与排查", 1)
add_table(
    ["问题现象", "可能原因", "解决方法"],
    [
        ["运行 .bat 提示 “未检测到 Python”", "未安装 Python 或未加入 PATH",
         "重新安装 Python 并勾选 Add Python to PATH"],
        ["pip install 失败 / 超时", "网络问题",
         "改用清华镜像：pip install -r requirements.txt -i https://pypi.tuna.tsinghua.edu.cn/simple"],
        ["启动 Web 报 ModuleNotFoundError", "依赖未装全",
         "在项目根目录重新执行 pip install -r requirements.txt"],
        ["页面打开但图表空白", "首次加载未完成 / 浏览器缓存",
         "等待 1 分钟后刷新；或 Ctrl + F5 强制刷新"],
        ["报 FileNotFoundError: scaler.pkl", "models/ 目录缺失模型文件",
         "运行 python train.py 重新训练并生成模型"],
        ["端口 5000 被占用", "其他程序占用端口",
         "关闭占用进程，或修改 web/app.py 末尾 app.run(port=5001)"],
        ["训练耗时过长", "数据量大 / 机器性能不足",
         "可在 config.py 中减小 RF_PARAMS / LGB_PARAMS 的 n_estimators"],
        ["中文乱码", "终端编码非 UTF-8",
         "脚本已加 chcp 65001；如手动运行请确认终端为 UTF-8"],
    ]
)

# ---------------- 8. 维护与关闭 ----------------
add_title("八、系统维护与关闭", 1)
add_title("8.1 关闭服务", 2)
add_bullet("在运行 web/app.py 的命令行窗口按 Ctrl + C 即可停止 Flask 服务；")
add_bullet("关闭 .bat 启动的命令行窗口同样会终止后台进程。")
add_title("8.2 重新训练", 2)
add_para("当数据集更新、特征逻辑改动或想调整模型参数时，需重新训练：")
add_code("python train.py")
add_para("训练完成后 models/ 下所有 .pkl 文件会被覆盖更新，无需手动清理。")
add_title("8.3 参数调整", 2)
add_para("所有超参数集中在 config.py，包括：RF_PARAMS、LR_PARAMS、LGB_PARAMS、"
         "ENSEMBLE_WEIGHTS（融合权重）、USEFULNESS_THRESHOLD（有用性阈值）、"
         "RANDOM_STATE、TEST_SIZE。修改后需重新运行 train.py 生效。")
add_title("8.4 日志与结果", 2)
add_bullet("训练评估报告：outputs/experiment_results.txt；")
add_bullet("Flask 运行日志：命令行窗口标准输出；")
add_bullet("如需保留日志，可使用 python web\\app.py > web.log 2>&1 重定向。")

# ---------------- 结尾 ----------------
doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
re = end.add_run("— 文档结束 —")
re.font.size = Pt(11)
re.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

doc.save(OUT)
print(f"已生成: {OUT}")
